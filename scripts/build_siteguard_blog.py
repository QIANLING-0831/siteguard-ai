from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "blog-assets"
OUT_DOCX = DOCS / "SiteGuard-AI-B现场巡检工作台-博客文章.docx"
OUT_MD = DOCS / "SiteGuard-AI-B现场巡检工作台-博客文章.md"

TITLE = "AI 工地安全巡检：从安全帽识别到证据闭环的产品设计实践"
SUBTITLE = "摄像头不是终点，识别也不是终点。真正困难的是把现场照片变成可复核、可派单、可追踪的安全管理流程。"

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "真正的难题，不是识别，而是流程",
        [
            "很多工地已经安装了摄像头，也积累了大量现场照片。可一旦把问题落到“谁没有佩戴安全帽”这件事上，管理人员仍然要在不同画面之间反复切换：先找人，再判断，再截图，再通知，最后还要确认整改是否完成。照片很多，真正能进入管理流程的证据却很少。",
            "这也是 B 现场巡检工作台想回答的问题：AI 不应该只在图片上画几个框，然后把结果扔给工作人员。它需要进入一条完整、可追溯的业务链路，把采集、识别、人工判断、隐患立案、整改和复核连接起来。",
        ],
    ),
    (
        "第一条原则：AI 发现，不等于隐患成立",
        [
            "在工作台里，模型输出被称为 Finding，也就是“发现”。它只是一个尚未被确认的异常观察。只有经过有权限的工作人员查看原图并确认后，它才会升级为 HazardCase，也就是需要跟踪处置的正式隐患。",
            "这个区分看起来像是命名问题，实际上决定了系统的责任边界。模型置信度只能说明模型对某个视觉判断的把握程度，不能替代安全员的专业判断，更不能直接代表风险等级。因此，AI 可以帮人排序、聚合和标记，但不能越过人工确认自动立案。",
        ],
    ),
    (
        "第二条原则：一张照片里的人，要逐个判断",
        [
            "真实工地画面通常不是“一个人、一张图”。一个楼层、材料场或脚手架区域里，可能同时出现十几名工作人员。系统首先检测画面中的人员区域，再对每个人分别判断安全帽状态，并把疑似未佩戴者关联到独立的 Finding。",
            "这里的“画面人员”只是当前证据中的临时检测区域，不是实名档案。原型不做人脸识别，也不跨照片追踪身份。这样既满足逐人审核的需要，也避免为了完成安全帽检测而不必要地扩大个人信息处理范围。",
        ],
    ),
    (
        "第三条原则：把人工精力留给不确定性",
        [
            "如果 AI 检测到几十个人，却要求安全员把每个人重新看一遍，系统只是把纸面工作搬到了屏幕上。B 工作台采用分级复核：高置信度的“已佩戴安全帽”通过初筛；高置信度的疑似违规进入独立抽检队列；处在阈值边缘、遮挡严重或画面模糊的结果，才优先进入人工审核。",
            "人工审核区提供单列、四宫格和九宫格三种视图。四宫格适合日常审核，九宫格适合快速扫视同一批次；每张卡片都保留人员框、Finding 编号、置信度、来源以及“确认隐患 / 驳回 AI”两个动作。进入派单、整改或复核后，由于信息和操作更复杂，界面再回到单列流程。",
        ],
    ),
    (
        "证据治理：重复照片不能靠“看起来像”就删除",
        [
            "连续抓拍很容易产生大量高度相似的照片。如果这些图片全部进入 AI 统计，同一个人可能被重复计数，同一个问题也可能生成多条 Finding。工作台因此把重复检测放在审核队列之前：疑似重复证据在人员作出决定前，不参与人员和 Finding 汇总。",
            "判断过程分为两层。第一层用项目、摄像头或采集来源以及五分钟时间窗口缩小候选范围；第二层再做图像对比。文件完全一致时使用 SHA-256 内容哈希，画面有轻微变化时使用感知哈希比较视觉结构。达到阈值后，系统只生成 DuplicateSuggestion，不会自动删除。",
            "工作人员展开重复图组后，可以按九宫格查看一张基准图和所有疑似重复图，逐张选择“保留”或“作废”。这里的作废也不是物理删除：原图、原因、操作人和时间都会保留在作废证据集合中；已经关联正式隐患或整改记录的证据，系统会阻止直接作废。",
        ],
    ),
    (
        "从本机摄像头到工地摄像头，接入方式需要分层",
        [
            "浏览器弹出的本机摄像头权限，只适合开发阶段验证采集、上传和识别链路。实际工地摄像头通常通过 RTSP、厂商协议或视频平台提供数据，浏览器无法直接稳定消费这些流。生产接入需要在摄像头与网页之间增加网关或媒体服务，把视频流转换为网页可用的预览、定时抓拍或事件快照。",
            "因此，网页负责项目切换、证据查看、人工审核和闭环操作；摄像头网关负责设备认证、断线重连、时间同步、抓拍策略和流媒体转换；视觉识别则放在独立适配器后面。即使模型运行环境暂时不可用，业务流程仍然可以使用模拟适配器完成演示和测试。",
        ],
    ),
    (
        "一条完整闭环，应该长什么样",
        [
            "01 采集：来自本机上传、固定摄像头或后续设备网关的原图进入同一个巡检批次。",
            "02 识别：模型逐图检测多人并生成 Finding，同时记录模型名称、适配器和置信度。",
            "03 分流：高置信戴帽通过初筛；高置信疑似结果进入抽检；不确定结果进入重点人工复核。",
            "04 确认：安全员查看原图和人员框，决定确认、驳回或暂缓。只有确认结果才能形成正式隐患。",
            "05 整改：隐患被指派给责任人并设置期限，责任人上传整改证据。",
            "06 复核：复核人比较整改前后证据，选择通过关闭或退回整改，整个过程写入审计记录。",
        ],
    ),
    (
        "原型之外，生产系统还要补哪些能力",
        [
            "B 页面目前仍是用于确定方向的交互原型。要进入真实项目，至少还需要补齐摄像头设备台账与健康监控、项目级权限控制、对象存储与证据保留策略、操作审计、模型版本管理、不同天气与光照条件下的评估，以及个人信息和数据安全合规。",
            "更重要的是，模型评估不能只看一个总准确率。安全帽大小、人员距离、遮挡、夜间照明、摄像头角度和画面压缩都会改变识别难度。生产评估应该按场景拆分，并持续记录误报、漏报和人工复核结果，用真实反馈校准阈值和审核策略。",
        ],
    ),
    (
        "结语：AI 的价值，是让责任链更清晰",
        [
            "工地安全巡检真正需要的，不是一个会画框的模型，而是一套能把“看见问题”变成“解决问题”的系统。AI 可以承担大规模初筛和证据整理，人负责确认事实、判断风险并承担最终责任；系统则负责让每一次判断都有原图、状态和操作记录可追溯。",
            "当识别、证据治理、人工审核和整改闭环被放在同一张工作台上，摄像头才不再只是被动录像设备，而成为现场安全管理流程中的一个可靠入口。",
        ],
    ),
]


def font(size: int, bold: bool = False):
    path = Path("C:/Windows/Fonts/msyh.ttc")
    fallback = Path("C:/Windows/Fonts/simhei.ttf")
    return ImageFont.truetype(str(path if path.exists() else fallback), size=size, index=0)


def crop_browser_assets() -> None:
    for source_name, output_name, max_height in [
        ("01-workbench-overview.png", "01-workbench-overview-crop.png", 620),
        ("03-review-grid.png", "03-review-grid-crop.png", 620),
    ]:
        image = Image.open(ASSETS / source_name).convert("RGB")
        image.crop((0, 0, image.width, min(max_height, image.height))).save(ASSETS / output_name, quality=94)


def build_duplicate_collage() -> None:
    width, height = 1800, 1780
    canvas = Image.new("RGB", (width, height), "#FBF7EF")
    draw = ImageDraw.Draw(canvas)
    title_font, subtitle_font = font(46, True), font(25)
    label_font, meta_font, small_font = font(24, True), font(20), font(18)
    draw.text((55, 38), "重复证据组 01｜九宫格比对", font=title_font, fill="#242B25")
    draw.text((55, 100), "以第一张为基准，系统只提示相似，工作人员逐张决定保留或作废", font=subtitle_font, fill="#6F6A61")

    source = Image.open(ROOT / "public" / "demo" / "site-floor.png").convert("RGB")
    margin, gap, top = 55, 22, 160
    card_width = (width - margin * 2 - gap * 2) // 3
    card_height = 500
    photo_height = 360
    times = ["00:05:43", "00:06:28", "00:06:43", "00:06:58", "00:07:13", "00:07:28", "00:07:43", "00:07:58", "00:08:13"]
    deltas = [0, 45, 60, 75, 90, 105, 120, 135, 150]

    for index in range(9):
        row, column = divmod(index, 3)
        x = margin + column * (card_width + gap)
        y = top + row * (card_height + gap)
        draw.rounded_rectangle((x, y, x + card_width, y + card_height), radius=24, fill="white", outline="#D9D1C2", width=3)
        photo = ImageOps.fit(source, (card_width - 6, photo_height), method=Image.Resampling.LANCZOS)
        canvas.paste(photo, (x + 3, y + 3))
        badge = "基准图" if index == 0 else "疑似重复"
        badge_color = "#2F7656" if index == 0 else "#9A5B28"
        badge_width = 110 if index == 0 else 150
        draw.rounded_rectangle((x + 18, y + 18, x + 18 + badge_width, y + 62), radius=20, fill=badge_color)
        draw.text((x + 34, y + 25), badge, font=small_font, fill="white")
        if index:
            draw.rounded_rectangle((x + card_width - 105, y + 18, x + card_width - 18, y + 62), radius=20, fill="#171B18")
            draw.text((x + card_width - 91, y + 25), "100%", font=small_font, fill="white")
        draw.text((x + 22, y + photo_height + 22), f"2号楼摄像头 08 · 连续抓拍 {index + 1:02d}", font=label_font, fill="#252A26")
        meta = f"{times[index]} · " + ("比对基准" if index == 0 else f"间隔 {deltas[index]} 秒 · 文件相同")
        draw.text((x + 22, y + photo_height + 65), meta, font=meta_font, fill="#756F65")
        action = "保留作为比对基准" if index == 0 else "保留 / 作废"
        action_color = "#2F7656" if index == 0 else "#8D542A"
        draw.text((x + 22, y + photo_height + 106), action, font=small_font, fill=action_color)

    canvas.save(ASSETS / "02-duplicate-grid-collage.png", quality=95)


def set_east_asia_font(run, east_asia: str = "Microsoft YaHei", latin: str = "Calibri") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, size: float, color: str, bold: bool = False) -> None:
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_east_asia_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    tail = paragraph.add_run(" 页")
    set_east_asia_font(tail)


def add_paragraph_shading(paragraph, fill: str, left_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "10")
        left.set(qn("w:color"), left_border)
        borders.append(left)
        p_pr.append(borders)


def set_image_alt(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 6.3) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    set_image_alt(shape, caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, "2A2F2B")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = doc.styles.add_style("Blog Title", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(title_style, 28, "203748", True)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(10)

    subtitle_style = doc.styles.add_style("Blog Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(subtitle_style, 14, "526677")
    subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_style.paragraph_format.space_after = Pt(22)
    subtitle_style.paragraph_format.line_spacing = 1.2

    caption = doc.styles["Caption"]
    set_style_font(caption, 9, "6F6A61")
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(12)
    caption.paragraph_format.keep_with_next = False

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("SiteGuard AI · 产品设计实践")
    set_east_asia_font(header_run)
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor.from_string("7B817C")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = doc.styles["Normal"]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string("7B817C")


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    for run in paragraph.runs:
        set_east_asia_font(run)


def add_workflow_steps(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        label, detail = text.split(" ", 1)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(7)
        label_run = paragraph.add_run(label + "  ")
        set_east_asia_font(label_run)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor.from_string("D45B2A")
        body_run = paragraph.add_run(detail)
        set_east_asia_font(body_run)


def add_closing_summary(doc: Document) -> None:
    paragraph = doc.add_paragraph("AI 负责规模化筛选，人负责事实确认与最终责任，系统负责证据和过程留痕。")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = paragraph.paragraph_format.right_indent = Inches(0.42)
    paragraph.paragraph_format.space_before = Pt(20)
    paragraph.paragraph_format.space_after = Pt(16)
    paragraph.paragraph_format.line_spacing = 1.25
    add_paragraph_shading(paragraph, "F4F6F9", "D45B2A")
    for run in paragraph.runs:
        set_east_asia_font(run)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor.from_string("203748")

    for label, text in [
        ("01", "识别结果先是 Finding，经过人员确认后才是正式隐患。"),
        ("02", "无效证据可以作废，但不能无痕删除。"),
        ("03", "效率来自分流和聚合，而不是取消人工判断。"),
    ]:
        line = doc.add_paragraph()
        line.paragraph_format.left_indent = Inches(0.72)
        line.paragraph_format.right_indent = Inches(0.5)
        line.paragraph_format.space_after = Pt(6)
        number = line.add_run(label + "  ")
        set_east_asia_font(number)
        number.bold = True
        number.font.color.rgb = RGBColor.from_string("D45B2A")
        body = line.add_run(text)
        set_east_asia_font(body)

    tags = doc.add_paragraph("关键词：AI 巡检 · 工地安全 · 安全帽识别 · 证据治理 · 人工复核 · 隐患闭环")
    tags.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tags.paragraph_format.space_before = Pt(14)
    for run in tags.runs:
        set_east_asia_font(run)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("7B817C")


def build_docx() -> None:
    doc = Document()
    configure_document(doc)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("PRODUCT PRACTICE  /  AI SAFETY")
    set_east_asia_font(run)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("D45B2A")

    doc.add_paragraph(TITLE, style="Blog Title")
    doc.add_paragraph(SUBTITLE, style="Blog Subtitle")
    meta = doc.add_paragraph("基于 SiteGuard AI · B 现场巡检工作台交互原型  |  2026 年 8 月")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(58)
    for run in meta.runs:
        set_east_asia_font(run)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string("7B817C")

    lead = doc.add_paragraph("真正可用的 AI 巡检，不是把“识别结果”堆给人，而是让机器负责大规模初筛，让人负责事实确认，让系统负责证据和责任链。")
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.5)
    lead.paragraph_format.right_indent = Inches(0.5)
    lead.paragraph_format.space_before = Pt(10)
    lead.paragraph_format.space_after = Pt(24)
    lead.paragraph_format.line_spacing = 1.25
    add_paragraph_shading(lead, "F4F6F9", "D45B2A")
    for run in lead.runs:
        set_east_asia_font(run)
        run.font.size = Pt(12.5)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("203748")

    note = doc.add_paragraph("说明：本文基于交互原型与合成演示数据。页面中的人数、Finding 数量和置信度仅用于展示工作流，不代表模型准确率或生产部署效果。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.left_indent = note.paragraph_format.right_indent = Inches(0.55)
    for run in note.runs:
        set_east_asia_font(run)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("756F65")

    doc.add_page_break()

    for index, (heading, paragraphs) in enumerate(SECTIONS):
        doc.add_heading(heading, level=1)
        if heading == "一条完整闭环，应该长什么样":
            add_workflow_steps(doc, paragraphs)
        else:
            for paragraph in paragraphs:
                add_body_paragraph(doc, paragraph)

        if heading == "结语：AI 的价值，是让责任链更清晰":
            add_closing_summary(doc)

        if index == 0:
            add_figure(doc, ASSETS / "01-workbench-overview-crop.png", "图 1｜B 现场巡检工作台总览：采集、AI 分流、人工审核与整改复核被放在同一条工作流中。")
        elif heading == "第三条原则：把人工精力留给不确定性":
            add_figure(doc, ASSETS / "03-review-grid-crop.png", "图 2｜人工审核九宫格：一屏浏览多条 Finding，并在缩略图卡片内直接确认或驳回。")
        elif heading == "证据治理：重复照片不能靠“看起来像”就删除":
            add_figure(doc, ASSETS / "02-duplicate-grid-collage.png", "图 3｜重复证据组九宫格：基准图与 8 张连续抓拍并排呈现，工作人员逐张决定保留或作废。", width=5.85)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "SiteGuard AI B 现场巡检工作台产品设计实践"
    doc.core_properties.author = "SiteGuard AI"
    doc.core_properties.keywords = "AI巡检, 工地安全, 安全帽识别, 证据治理, 人工审核, 隐患闭环"
    doc.save(OUT_DOCX)


def build_markdown() -> None:
    parts = [
        f"# {TITLE}",
        "",
        f"> {SUBTITLE}",
        "",
        "**说明：** 本文基于交互原型与合成演示数据。页面中的人数、Finding 数量和置信度仅用于展示工作流，不代表模型准确率或生产部署效果。",
        "",
    ]
    for index, (heading, paragraphs) in enumerate(SECTIONS):
        parts.extend([f"## {heading}", ""])
        parts.extend([paragraph + "\n" for paragraph in paragraphs])
        if index == 0:
            parts.extend(["![B 现场巡检工作台总览](./blog-assets/01-workbench-overview-crop.png)", "", "*图 1｜B 现场巡检工作台总览：采集、AI 分流、人工审核与整改复核被放在同一条工作流中。*", ""])
        elif heading == "第三条原则：把人工精力留给不确定性":
            parts.extend(["![人工审核九宫格](./blog-assets/03-review-grid-crop.png)", "", "*图 2｜人工审核九宫格：一屏浏览多条 Finding，并在缩略图卡片内直接确认或驳回。*", ""])
        elif heading == "证据治理：重复照片不能靠“看起来像”就删除":
            parts.extend(["![重复证据组九宫格](./blog-assets/02-duplicate-grid-collage.png)", "", "*图 3｜重复证据组九宫格：基准图与 8 张连续抓拍并排呈现，工作人员逐张决定保留或作废。*", ""])
        elif heading == "结语：AI 的价值，是让责任链更清晰":
            parts.extend(["> AI 负责规模化筛选，人负责事实确认与最终责任，系统负责证据和过程留痕。", "", "**三个判断标准：** 识别结果先是 Finding，确认后才是隐患；无效证据可以作废，但不能无痕删除；效率来自分流和聚合，而不是取消人工判断。", ""])
    OUT_MD.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")


if __name__ == "__main__":
    ASSETS.mkdir(parents=True, exist_ok=True)
    crop_browser_assets()
    build_duplicate_collage()
    build_markdown()
    build_docx()
    print(OUT_MD)
    print(OUT_DOCX)
